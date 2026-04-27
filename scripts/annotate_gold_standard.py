"""CLI helper for building and validating the LUCA gold standard.

The gold standard is the manual ground truth against which the L3
structured extractor (and, later, the LLM baseline) is evaluated. It
**must** be filled in by hand by reading the official 10-K filings; this
script only manages the *structural* skeleton of the workbook and never
pre-fills numeric values.

Subcommands
-----------
``init``
    Create ``data/gold_standard.xlsx`` with a header row and one empty
    row per (ticker, metric) pair. Refuses to overwrite an existing
    file.

``validate``
    Read the workbook, check that every row is fully and consistently
    annotated, and print a per-ticker summary. Read-only.

``status``
    Print a short progress report (``X/N rows annotated``) and list the
    (ticker, metric) pairs that are still missing. Read-only.

Usage (always invoke through the project's venv interpreter):

    /home/<user>/luca-ai-lab/venv/bin/python \
        scripts/annotate_gold_standard.py init
    /home/<user>/luca-ai-lab/venv/bin/python \
        scripts/annotate_gold_standard.py validate
    /home/<user>/luca-ai-lab/venv/bin/python \
        scripts/annotate_gold_standard.py status

Reglas duras (no negociables):
    * Nunca pre-rellenamos ``value``, ``fiscal_year``, ``unit``,
      ``source_page`` ni ``notes``. La regla de oro es que los datos los
      escribe el usuario a mano leyendo el 10-K. Pre-rellenar invalida
      el experimento del paper.
    * El script jamás llama a EDGAR ni a un LLM. Es estrictamente
      offline.
"""

from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# When invoked as a script (``python scripts/annotate_gold_standard.py``)
# the project root is not on ``sys.path``. Add it before the package import
# so the constants module resolves cleanly without forcing the user to set
# ``PYTHONPATH``.
_PROJECT_ROOT: Path = Path(__file__).resolve().parent.parent
if str(_PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(_PROJECT_ROOT))

from agents._gold_schema import (  # noqa: E402  (path tweak above)
    GOLD_COLUMNS,
    GOLD_METRICS,
    GOLD_ROW_COUNT,
    GOLD_TICKERS,
)

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

#: Default location of the workbook relative to the project root.
DEFAULT_WORKBOOK_PATH: Path = _PROJECT_ROOT / "data" / "gold_standard.xlsx"

#: Excel sheet name for the annotations.
SHEET_NAME: str = "gold_standard"

#: Min/max sane values for ``fiscal_year`` (4-digit year between 1990 and
#: 2099 inclusive). Anything outside this range is almost certainly a typo.
_FY_MIN: int = 1990
_FY_MAX: int = 2099


# ---------------------------------------------------------------------------
# Validation result types
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class RowIssue:
    """A single annotation problem found during ``validate``.

    Attributes:
        excel_row: 1-based Excel row number (header is row 1, data starts
            at row 2).
        ticker: Ticker for that row, or ``None`` if the cell is empty.
        metric: Metric name for that row, or ``None`` if the cell is empty.
        reason: Human-readable Spanish description of the problem.
    """

    excel_row: int
    ticker: str | None
    metric: str | None
    reason: str


@dataclass(frozen=True)
class ValidationReport:
    """Aggregated outcome of a ``validate`` run."""

    total_rows: int
    annotated_rows: int
    empty_rows: int
    invalid_rows: int
    issues: tuple[RowIssue, ...]

    @property
    def is_clean(self) -> bool:
        """True when there are zero issues."""
        return not self.issues


# ---------------------------------------------------------------------------
# Workbook initialisation
# ---------------------------------------------------------------------------


def _build_skeleton_rows() -> list[tuple[str | None, ...]]:
    """Build the empty data rows: one per (ticker, metric) pair.

    Only ``ticker`` and ``metric`` are pre-filled; every other column is
    left as ``None`` so openpyxl writes a truly empty cell. Order:
    tickers in :data:`GOLD_TICKERS` order, metrics in
    :data:`GOLD_METRICS` order within each ticker.
    """
    rows: list[tuple[str | None, ...]] = []
    for ticker in GOLD_TICKERS:
        for metric in GOLD_METRICS:
            row: list[str | None] = [None] * len(GOLD_COLUMNS)
            row[GOLD_COLUMNS.index("ticker")] = ticker
            row[GOLD_COLUMNS.index("metric")] = metric
            rows.append(tuple(row))
    return rows


def _style_header(sheet: Worksheet) -> None:
    """Apply lightweight styling to the header row.

    Bold font + light grey fill make it obvious where the header ends and
    where the user starts typing. Column widths are widened to fit the
    typical content (numeric values up to trillions, plus URLs/notes).
    """
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
    for idx, _ in enumerate(GOLD_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=idx)
        cell.font = header_font
        cell.fill = header_fill

    # Reasonable default widths in Excel character units.
    widths = {
        "ticker": 10,
        "fiscal_year": 14,
        "metric": 20,
        "value": 18,
        "unit": 10,
        "source_page": 14,
        "notes": 40,
    }
    for col_idx, name in enumerate(GOLD_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(col_idx)].width = widths.get(name, 14)


def init_workbook(path: Path) -> bool:
    """Create the gold-standard workbook at ``path``.

    Args:
        path: Destination ``.xlsx`` path. Parent directories are created
            if they don't exist yet.

    Returns:
        ``True`` if a new file was written, ``False`` if a file already
        existed at ``path`` (in which case it is left untouched).
    """
    if path.exists():
        return False

    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    sheet = wb.active
    sheet.title = SHEET_NAME

    sheet.append(list(GOLD_COLUMNS))
    for row in _build_skeleton_rows():
        sheet.append(list(row))

    _style_header(sheet)

    # Freeze the header so it stays visible while the user scrolls.
    sheet.freeze_panes = "A2"

    wb.save(path)
    return True


# ---------------------------------------------------------------------------
# Workbook reading + validation
# ---------------------------------------------------------------------------


def _open_sheet(path: Path) -> Worksheet:
    """Open the annotation workbook and return its annotation sheet.

    Raises:
        FileNotFoundError: The workbook does not exist.
        ValueError: The expected sheet or header is missing/incorrect.
    """
    if not path.exists():
        raise FileNotFoundError(
            f"No se encuentra el archivo de gold standard en {path}. "
            "Ejecuta 'init' primero."
        )

    wb = load_workbook(path, data_only=True, read_only=False)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(
            f"El archivo {path} no contiene la hoja '{SHEET_NAME}'. "
            "Posiblemente fue renombrada o sustituida manualmente."
        )

    sheet = wb[SHEET_NAME]
    header = tuple(
        (sheet.cell(row=1, column=idx + 1).value or "")
        for idx in range(len(GOLD_COLUMNS))
    )
    if header != GOLD_COLUMNS:
        raise ValueError(
            f"Cabecera inesperada en {path}: {header!r}. "
            f"Se esperaba {GOLD_COLUMNS!r}."
        )
    return sheet


def _is_blank(value: object) -> bool:
    """Return True for cell values that count as empty."""
    if value is None:
        return True
    if isinstance(value, str) and not value.strip():
        return True
    return False


def _row_is_fully_blank(values: Sequence[object]) -> bool:
    """A row is "blank" only if *every* annotated column is empty."""
    return all(_is_blank(v) for v in values)


def _validate_row(
    excel_row: int,
    raw: Sequence[object],
) -> tuple[bool, list[RowIssue]]:
    """Validate a single non-skeleton data row.

    Args:
        excel_row: 1-based Excel row index for issue reporting.
        raw: Cell values in :data:`GOLD_COLUMNS` order.

    Returns:
        A tuple ``(annotated, issues)``. ``annotated`` is ``True`` when
        at least one of the user-fillable columns has content (so the
        row is treated as a real attempted annotation rather than the
        untouched skeleton). ``issues`` lists every problem found; an
        empty list means the row is valid.
    """
    by_name: dict[str, object] = dict(zip(GOLD_COLUMNS, raw))

    user_columns = ("fiscal_year", "value", "unit", "source_page", "notes")
    user_filled = any(not _is_blank(by_name[c]) for c in user_columns)

    issues: list[RowIssue] = []
    ticker = by_name["ticker"] if isinstance(by_name["ticker"], str) else None
    metric = by_name["metric"] if isinstance(by_name["metric"], str) else None

    def add(reason: str) -> None:
        issues.append(
            RowIssue(
                excel_row=excel_row,
                ticker=ticker,
                metric=metric,
                reason=reason,
            )
        )

    # Skeleton columns must always be present (even on untouched rows).
    if _is_blank(by_name["ticker"]):
        add("falta 'ticker'")
    elif ticker not in GOLD_TICKERS:
        add(f"ticker desconocido: {by_name['ticker']!r}")

    if _is_blank(by_name["metric"]):
        add("falta 'metric'")
    elif metric not in GOLD_METRICS:
        add(f"metric desconocida: {by_name['metric']!r}")

    if not user_filled:
        # Untouched skeleton row: the only "issue" is being incomplete,
        # which we report as a single problem so callers can count it.
        add("fila sin anotar")
        return False, issues

    # --- fiscal_year: int 4-digit, in sane range -------------------------
    fy = by_name["fiscal_year"]
    if _is_blank(fy):
        add("falta 'fiscal_year'")
    else:
        fy_int = _coerce_int(fy)
        if fy_int is None:
            add(f"'fiscal_year' no es entero: {fy!r}")
        elif not (_FY_MIN <= fy_int <= _FY_MAX):
            add(f"'fiscal_year' fuera de rango razonable: {fy_int}")

    # --- value: numeric (int or float) -----------------------------------
    val = by_name["value"]
    if _is_blank(val):
        add("falta 'value'")
    elif not isinstance(val, (int, float)) or isinstance(val, bool):
        # bool is a subclass of int — explicitly reject it.
        add(f"'value' no es numérico: {val!r}")

    # --- unit: non-empty string ------------------------------------------
    unit = by_name["unit"]
    if _is_blank(unit):
        add("falta 'unit'")
    elif not isinstance(unit, str):
        add(f"'unit' no es texto: {unit!r}")

    # --- source_page: int -------------------------------------------------
    sp = by_name["source_page"]
    if _is_blank(sp):
        add("falta 'source_page'")
    else:
        sp_int = _coerce_int(sp)
        if sp_int is None:
            add(f"'source_page' no es entero: {sp!r}")
        elif sp_int < 1:
            add(f"'source_page' debe ser >= 1: {sp_int}")

    # 'notes' is optional — not validated.
    return True, issues


def _coerce_int(value: object) -> int | None:
    """Best-effort conversion of an Excel cell value to ``int``.

    Excel often stores integers as floats (``2024.0``); we accept those
    silently. Strings of pure digits also work, to be lenient with users
    who paste from PDFs. Anything else returns ``None``.
    """
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if value.is_integer():
            return int(value)
        return None
    if isinstance(value, str):
        stripped = value.strip()
        if stripped.isdigit():
            return int(stripped)
        try:
            f = float(stripped)
        except ValueError:
            return None
        if f.is_integer():
            return int(f)
        return None
    return None


def validate_workbook(path: Path) -> ValidationReport:
    """Read the workbook and produce a :class:`ValidationReport`."""
    sheet = _open_sheet(path)

    issues: list[RowIssue] = []
    annotated = 0
    empty = 0
    invalid = 0
    total = 0

    for excel_row in range(2, sheet.max_row + 1):
        raw = tuple(
            sheet.cell(row=excel_row, column=idx + 1).value
            for idx in range(len(GOLD_COLUMNS))
        )
        if _row_is_fully_blank(raw):
            # Truly empty trailing rows can appear after edits in Excel;
            # ignore them silently rather than counting as missing data.
            continue

        total += 1
        is_annotated, row_issues = _validate_row(excel_row, raw)
        if is_annotated:
            annotated += 1
            if row_issues:
                invalid += 1
                issues.extend(row_issues)
        else:
            empty += 1
            issues.extend(row_issues)

    return ValidationReport(
        total_rows=total,
        annotated_rows=annotated,
        empty_rows=empty,
        invalid_rows=invalid,
        issues=tuple(issues),
    )


# ---------------------------------------------------------------------------
# Status (lightweight, read-only)
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class StatusReport:
    """Lightweight progress report for the ``status`` subcommand."""

    total_rows: int
    annotated_rows: int
    missing_pairs: tuple[tuple[str, str], ...]


def status_workbook(path: Path) -> StatusReport:
    """Compute a quick progress summary without reporting per-row errors."""
    sheet = _open_sheet(path)

    annotated_pairs: set[tuple[str, str]] = set()
    total = 0

    for excel_row in range(2, sheet.max_row + 1):
        raw = tuple(
            sheet.cell(row=excel_row, column=idx + 1).value
            for idx in range(len(GOLD_COLUMNS))
        )
        if _row_is_fully_blank(raw):
            continue

        total += 1
        by_name = dict(zip(GOLD_COLUMNS, raw))
        ticker = by_name["ticker"]
        metric = by_name["metric"]
        user_columns = ("fiscal_year", "value", "unit", "source_page")
        if all(not _is_blank(by_name[c]) for c in user_columns):
            if isinstance(ticker, str) and isinstance(metric, str):
                annotated_pairs.add((ticker, metric))

    expected_pairs = {(t, m) for t in GOLD_TICKERS for m in GOLD_METRICS}
    missing = sorted(expected_pairs - annotated_pairs)

    return StatusReport(
        total_rows=total,
        annotated_rows=len(annotated_pairs),
        missing_pairs=tuple(missing),
    )


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def _parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        prog="annotate_gold_standard.py",
        description=(
            "Gestiona el Excel de gold standard de LUCA "
            "(crear esqueleto, validar anotaciones, ver progreso)."
        ),
    )
    parser.add_argument(
        "--path",
        type=Path,
        default=DEFAULT_WORKBOOK_PATH,
        help=(
            f"Ruta del archivo Excel "
            f"(por defecto: {DEFAULT_WORKBOOK_PATH.relative_to(_PROJECT_ROOT)})."
        ),
    )

    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser(
        "init",
        help="Crea el Excel con el esqueleto (ticker, metric) si no existe.",
    )
    sub.add_parser(
        "validate",
        help="Valida que todas las filas estén anotadas con tipos sanos.",
    )
    sub.add_parser(
        "status",
        help="Imprime el progreso de anotación.",
    )
    return parser.parse_args(argv)


def _print_init(path: Path, created: bool) -> None:
    if created:
        print(f"Esqueleto creado en {path}")
        print(
            f"  {len(GOLD_TICKERS)} tickers x {len(GOLD_METRICS)} métricas "
            f"= {GOLD_ROW_COUNT} filas listas para anotar."
        )
        print(
            "Recuerda: rellena 'value', 'fiscal_year', 'unit' y 'source_page' "
            "leyendo el 10-K oficial. NUNCA copies números de un sistema automático."
        )
    else:
        print(f"El archivo {path} ya existe; no se sobrescribe.")
        print("Borra o renombra el archivo si quieres regenerar el esqueleto.")


def _print_validate(report: ValidationReport) -> None:
    print(
        f"Filas detectadas: {report.total_rows}  "
        f"(anotadas: {report.annotated_rows}, vacías: {report.empty_rows}, "
        f"inválidas: {report.invalid_rows})"
    )
    if report.is_clean:
        print("Todas las filas están completas y son válidas.")
        return

    print("Problemas encontrados:")
    for issue in report.issues:
        ident = f"{issue.ticker or '?'} / {issue.metric or '?'}"
        print(f"  fila {issue.excel_row:>3} [{ident}]: {issue.reason}")


def _print_status(report: StatusReport) -> None:
    print(
        f"{report.annotated_rows}/{GOLD_ROW_COUNT} filas anotadas "
        f"({report.total_rows} filas con contenido)."
    )
    if not report.missing_pairs:
        print("No queda ninguna combinación (ticker, métrica) por anotar.")
        return

    print("Combinaciones pendientes:")
    # Group by ticker so the user can plan ticker-by-ticker.
    by_ticker: dict[str, list[str]] = {}
    for ticker, metric in report.missing_pairs:
        by_ticker.setdefault(ticker, []).append(metric)
    for ticker in GOLD_TICKERS:
        if ticker in by_ticker:
            metrics = ", ".join(by_ticker[ticker])
            print(f"  {ticker}: {metrics}")


def main(argv: list[str] | None = None) -> int:
    """CLI entry point. Returns a process exit code."""
    args = _parse_args(argv)
    path: Path = args.path

    if args.command == "init":
        created = init_workbook(path)
        _print_init(path, created)
        return 0

    if args.command == "validate":
        try:
            report = validate_workbook(path)
        except (FileNotFoundError, ValueError) as exc:
            print(f"Error: {exc}", file=sys.stderr)
            return 2
        _print_validate(report)
        return 0 if report.is_clean else 1

    if args.command == "status":
        try:
            report = status_workbook(path)
        except (FileNotFoundError, ValueError) as exc:
            print(f"Error: {exc}", file=sys.stderr)
            return 2
        _print_status(report)
        return 0

    # argparse with required=True should make this unreachable.
    print(f"Subcomando desconocido: {args.command}", file=sys.stderr)
    return 2


if __name__ == "__main__":
    raise SystemExit(main())


__all__ = [
    "DEFAULT_WORKBOOK_PATH",
    "RowIssue",
    "SHEET_NAME",
    "StatusReport",
    "ValidationReport",
    "init_workbook",
    "main",
    "status_workbook",
    "validate_workbook",
]
