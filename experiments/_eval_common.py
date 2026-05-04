"""Shared helpers for the Sprint 1 evaluation experiments.

This module is consumed by:

* :mod:`experiments.eval_extraction` (T1.8) — compares the L3 structured
  extractor against the manual gold standard.
* :mod:`experiments.eval_baseline_llm` (T1.9a / T1.9b) — compares an
  LLM-only baseline against the same gold standard.

Both experiments need the *same* gold-standard reader and the *same*
numeric comparator, so centralising the logic here is the only way to
guarantee that L3 and the LLM baseline are evaluated against an
identical ground-truth interpretation. Re-implementing
``coerce_gold_value`` or the comparison rules locally would silently
drift over time.

Reading the gold standard
-------------------------
The workbook lives at ``data/gold_standard.xlsx`` (60 rows, one per
``(ticker, metric)`` pair). The ``value`` column is either:

* a numeric (``int``/``float``) for present metrics, or
* the case-insensitive sentinel ``"null"`` for metrics that are
  *structurally* absent from a filer (e.g. JPM does not publish
  ``OperatingIncomeLoss``).

We delegate the coercion of those two cases to
:func:`agents._gold_schema.coerce_gold_value`. A failure to coerce a
single row is fatal: this evaluator must not silently skip rows.

Comparing values
----------------
:func:`compare_values` implements a four-case decision table:

* ``gold is None`` and ``predicted is None`` → match. Both agree the
  metric is absent.
* ``gold is None`` and ``predicted is not None`` → mismatch. The
  extractor hallucinated a value.
* ``gold is not None`` and ``predicted is None`` → mismatch. The
  extractor missed a known value.
* both numeric → match when ``abs(g-p)/abs(g) < rel_tol``. The default
  tolerance is 0.5%; see the function docstring for the reasoning.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Final

from openpyxl import load_workbook

from agents._gold_schema import (
    GOLD_COLUMNS,
    GOLD_METRICS,
    GOLD_TICKERS,
    coerce_gold_value,
)

#: Worksheet name produced by ``scripts/annotate_gold_standard.py``.
GOLD_SHEET_NAME: Final[str] = "gold_standard"


@dataclass(frozen=True)
class GoldEntry:
    """A single annotated row of the gold standard.

    Attributes:
        value: The annotated metric value, already coerced. ``None``
            means the metric is structurally absent for that filer.
        fiscal_year: 4-digit fiscal year as annotated (or ``None`` if
            the cell was left blank — should not happen on a clean
            workbook, but the type is permissive to keep the reader
            tolerant).
        unit: Unit string as annotated (e.g. ``"USD"`` or
            ``"USD/share"``). Kept verbatim from the workbook; we do
            not enforce ``EXPECTED_UNITS`` here because the dedicated
            ``annotate_gold_standard.py validate`` subcommand already
            checks that.
        source_page: Either an integer page number or a free-text
            statement name (e.g. ``"CONSOLIDATED STATEMENTS OF
            OPERATIONS"``). Metadata only — never compared.
        notes: Optional free-text annotation notes.
    """

    value: float | None
    fiscal_year: int | None
    unit: str
    source_page: str | int
    notes: str


def load_gold_standard(path: Path) -> dict[tuple[str, str], GoldEntry]:
    """Load the gold-standard workbook into a ``(ticker, metric) → GoldEntry`` map.

    Args:
        path: Path to the Excel workbook produced (and validated) by
            ``scripts/annotate_gold_standard.py``.

    Returns:
        A dict keyed by ``(ticker, metric)`` tuples in the canonical
        :data:`GOLD_TICKERS` × :data:`GOLD_METRICS` ordering.

    Raises:
        FileNotFoundError: When ``path`` does not exist.
        ValueError: When the workbook is structurally wrong (missing
            sheet, wrong header) or when any row fails coercion. The
            error message identifies the offending row by Excel index
            so the annotator can fix it manually.
    """
    if not path.is_file():
        raise FileNotFoundError(
            f"No se encuentra el gold standard en {path}. "
            "Genera el archivo con scripts/annotate_gold_standard.py init."
        )

    wb = load_workbook(path, data_only=True, read_only=True)
    if GOLD_SHEET_NAME not in wb.sheetnames:
        raise ValueError(
            f"El archivo {path} no contiene la hoja {GOLD_SHEET_NAME!r}."
        )

    sheet = wb[GOLD_SHEET_NAME]
    rows_iter = sheet.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if header is None or tuple(c or "" for c in header[: len(GOLD_COLUMNS)]) != GOLD_COLUMNS:
        raise ValueError(
            f"Cabecera inesperada en {path}: {header!r}. "
            f"Se esperaba {GOLD_COLUMNS!r}."
        )

    entries: dict[tuple[str, str], GoldEntry] = {}
    valid_tickers = set(GOLD_TICKERS)
    valid_metrics = set(GOLD_METRICS)

    # Excel header is row 1; data starts at row 2. Track index for error
    # messages.
    for offset, row in enumerate(rows_iter, start=2):
        # Skip purely-empty trailing rows (Excel sometimes leaves them).
        if row is None or all(cell is None or (isinstance(cell, str) and not cell.strip()) for cell in row):
            continue

        # Pad the row to GOLD_COLUMNS length defensively.
        padded = list(row) + [None] * (len(GOLD_COLUMNS) - len(row))
        by_name = dict(zip(GOLD_COLUMNS, padded))

        ticker = by_name["ticker"]
        metric = by_name["metric"]
        if ticker not in valid_tickers or metric not in valid_metrics:
            raise ValueError(
                f"Fila {offset} en {path}: combinación desconocida "
                f"ticker={ticker!r}, metric={metric!r}."
            )

        try:
            value = coerce_gold_value(by_name["value"])
        except ValueError as exc:
            raise ValueError(
                f"Fila {offset} en {path} (ticker={ticker}, "
                f"metric={metric}): {exc}"
            ) from exc

        fiscal_year_raw = by_name["fiscal_year"]
        if isinstance(fiscal_year_raw, int) and not isinstance(fiscal_year_raw, bool):
            fiscal_year: int | None = fiscal_year_raw
        elif isinstance(fiscal_year_raw, float) and fiscal_year_raw.is_integer():
            fiscal_year = int(fiscal_year_raw)
        else:
            fiscal_year = None

        unit_raw = by_name["unit"]
        unit = unit_raw if isinstance(unit_raw, str) else ""

        source_page_raw = by_name["source_page"]
        source_page: str | int
        if isinstance(source_page_raw, int) and not isinstance(source_page_raw, bool):
            source_page = source_page_raw
        elif isinstance(source_page_raw, float) and source_page_raw.is_integer():
            source_page = int(source_page_raw)
        elif isinstance(source_page_raw, str):
            source_page = source_page_raw
        else:
            source_page = ""

        notes_raw = by_name["notes"]
        notes = notes_raw if isinstance(notes_raw, str) else ""

        entries[(ticker, metric)] = GoldEntry(
            value=value,
            fiscal_year=fiscal_year,
            unit=unit,
            source_page=source_page,
            notes=notes,
        )

    return entries


def compare_values(
    gold: float | None,
    predicted: float | None,
    *,
    rel_tol: float = 0.005,
) -> tuple[bool, float | None, float | None]:
    """Compare a gold value against a predicted value.

    The decision table is:

    +-----------------+--------------------+----------------------------------+
    | gold            | predicted          | result                           |
    +=================+====================+==================================+
    | ``None``        | ``None``           | ``(True, None, None)``           |
    +-----------------+--------------------+----------------------------------+
    | ``None``        | numeric            | ``(False, None, None)``          |
    +-----------------+--------------------+----------------------------------+
    | numeric         | ``None``           | ``(False, abs(gold), None)``     |
    +-----------------+--------------------+----------------------------------+
    | numeric         | numeric            | match iff                        |
    |                 |                    | ``abs(g-p) / abs(g) < rel_tol``  |
    +-----------------+--------------------+----------------------------------+

    On the default tolerance (0.5%):
        Real-world rounding noise between the gold standard (where the
        annotator copies the values printed in the 10-K, typically with
        the millions/billions rounding chosen by the issuer) and XBRL
        (which carries the unrounded figure) is consistently below
        0.1%. A 0.5% bound therefore absorbs *any* rounding-induced
        difference while still being tight enough to flag genuine
        extraction errors (which are usually orders-of-magnitude wrong
        — wrong fiscal year, wrong tag, wrong scaling — not 0.5%
        wrong). T1.10 may revisit this number once we have the per-row
        error distribution.

    Args:
        gold: The annotated value (already coerced — ``None`` means
            "structurally absent").
        predicted: The extractor's prediction (``None`` means "no
            value found").
        rel_tol: Relative tolerance for the numeric-vs-numeric case.
            Defaults to 0.5%.

    Returns:
        Tuple ``(match, abs_error, rel_error)``. ``abs_error`` and
        ``rel_error`` are ``None`` whenever the comparison cannot
        produce a numeric error (any branch where at least one side is
        ``None`` collapses one or both error components).
    """
    if gold is None and predicted is None:
        return True, None, None
    if gold is None and predicted is not None:
        # Hallucinated value — no abs/rel error is meaningful here.
        return False, None, None
    if gold is not None and predicted is None:
        return False, abs(gold), None

    # Both numeric.
    assert gold is not None and predicted is not None  # for type checkers
    abs_error = abs(gold - predicted)
    if gold == 0.0:
        # Pathological: the gold value itself is zero. We cannot compute
        # a relative error, so treat it as match iff the predicted is
        # also exactly zero. Not expected in our gold standard (no
        # legitimate fiscal-year metric is zero) but covered for
        # robustness.
        return abs_error == 0.0, abs_error, None
    rel_error = abs_error / abs(gold)
    return rel_error < rel_tol, abs_error, rel_error


__all__ = [
    "GOLD_SHEET_NAME",
    "GoldEntry",
    "compare_values",
    "load_gold_standard",
]
